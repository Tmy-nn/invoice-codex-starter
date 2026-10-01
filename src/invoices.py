"""Validate a monthly workbook and produce sample PDFs and unsent email files."""

from __future__ import annotations

import csv
import json
import re
import uuid
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from email.message import EmailMessage
from pathlib import Path

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas


HEADERS = ("customer_id", "customer_name", "email", "description", "amount_yen")
MONTH_RE = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
ID_RE = re.compile(r"^[A-Za-z0-9_-]{1,40}$")
EMAIL_RE = re.compile(r"^[^\s@<>;,]+@[^\s@<>;,]+\.[^\s@<>;,]+$")
INK = colors.HexColor("#26313D")
MUTED = colors.HexColor("#52606D")
LINE = colors.HexColor("#D8DEE4")
PALE = colors.HexColor("#F1F3F5")
FONT = "MPLUS1p"
FONT_FILE = Path(__file__).resolve().parent / "fonts" / "MPLUS1p-Regular.ttf"


class InvoiceError(ValueError):
    """A validation or safe-output failure."""


class InvoiceRunError(RuntimeError):
    """A batch failed after its output directory was created."""

    def __init__(self, run_dir: Path):
        self.run_dir = run_dir
        super().__init__(f"batch incomplete: {run_dir}")


@dataclass(frozen=True)
class InvoiceRow:
    customer_id: str
    customer_name: str
    email: str
    description: str
    amount_yen: int


def _cell_text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def validate_month(month: str) -> str:
    if not MONTH_RE.fullmatch(month):
        raise InvoiceError("month must be YYYY-MM")
    return month


def read_rows(path: Path) -> list[InvoiceRow]:
    if not path.is_file():
        raise InvoiceError(f"input file is missing: {path}")
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.active
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None or tuple(_cell_text(x) for x in header[:5]) != HEADERS:
            raise InvoiceError(f"first five columns must be: {', '.join(HEADERS)}")
        result: list[InvoiceRow] = []
        seen: set[str] = set()
        errors: list[str] = []
        for line_no, row in enumerate(rows, start=2):
            values = list(row[:5])
            values += [None] * (5 - len(values))
            if not any(_cell_text(v) for v in values):
                continue
            customer_id, name, email, description = map(_cell_text, values[:4])
            amount_raw = values[4]
            problems: list[str] = []
            if not ID_RE.fullmatch(customer_id):
                problems.append("invalid customer_id")
            elif customer_id in seen:
                problems.append("duplicate customer_id")
            if not name or len(name) > 60 or name[0] in "=+-@\t" or any(c in name for c in "\r\n"):
                problems.append("invalid customer_name")
            if not EMAIL_RE.fullmatch(email):
                problems.append("invalid email")
            if not description or len(description) > 65 or any(c in description for c in "\r\n"):
                problems.append("invalid description")
            try:
                amount = Decimal(str(amount_raw))
                if not amount.is_finite() or amount != amount.to_integral_value() or not 1 <= amount <= 100_000_000:
                    problems.append("amount_yen must be a positive integer up to 100000000")
            except (InvalidOperation, TypeError, ValueError):
                amount = Decimal(0)
                problems.append("invalid amount_yen")
            if problems:
                errors.append(f"row {line_no}: {', '.join(problems)}")
            else:
                seen.add(customer_id)
                result.append(InvoiceRow(customer_id, name, email, description, int(amount)))
        if errors:
            raise InvoiceError("workbook rejected; no files generated:\n" + "\n".join(errors[:30]))
        if not result:
            raise InvoiceError("workbook has no invoice rows")
        return result
    finally:
        book.close()


def _font() -> None:
    if FONT not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(FONT, str(FONT_FILE)))


def validate_pdf_fit(rows: list[InvoiceRow], month: str) -> None:
    """Reject text that would collide with the other invoice column."""
    _font()
    printable = A4[0] - 40 * mm
    month_width = pdfmetrics.stringWidth(f"請求月  {month}", FONT, 10.5)
    problems = []
    for index, row in enumerate(rows, start=2):
        name_width = pdfmetrics.stringWidth(row.customer_name + " 御中", FONT, 13)
        amount_width = pdfmetrics.stringWidth(f"¥{row.amount_yen:,}", FONT, 10.5)
        desc_width = pdfmetrics.stringWidth(row.description, FONT, 10.5)
        if name_width + month_width + 8 * mm > printable:
            problems.append(f"row {index}: customer_name does not fit PDF")
        if desc_width + amount_width + 16 * mm > printable:
            problems.append(f"row {index}: description does not fit PDF")
    if problems:
        raise InvoiceError("PDF text width rejected; no files generated:\n" + "\n".join(problems[:30]))


def _text(canvas: Canvas, x: float, y: float, value: str, size: float = 9, color=INK) -> None:
    canvas.setFillColor(color)
    canvas.setFont(FONT, size)
    canvas.drawString(x, y, value)


def _right(canvas: Canvas, x: float, y: float, value: str, size: float = 9, color=INK) -> None:
    canvas.setFillColor(color)
    canvas.setFont(FONT, size)
    canvas.drawRightString(x, y, value)


def render_pdf(path: Path, month: str, row: InvoiceRow) -> None:
    """Create an A4 sample invoice. No tax breakdown or real sender is implied."""
    _font()
    page_w, page_h = A4
    left, right = 20 * mm, page_w - 20 * mm
    canvas = Canvas(str(path), pagesize=A4, pageCompression=1)
    canvas.setTitle(f"サンプル請求書 {month} {row.customer_id}")
    _text(canvas, left, page_h - 25 * mm, "請 求 書", 17)
    _right(canvas, right, page_h - 25 * mm, "SAMPLE / 架空データ", 10.5, MUTED)
    canvas.setStrokeColor(LINE)
    canvas.line(left, page_h - 30 * mm, right, page_h - 30 * mm)

    y = page_h - 45 * mm
    _text(canvas, left, y, f"{row.customer_name} 御中", 13)
    _right(canvas, right, y, f"請求月  {month}", 10.5)
    y -= 18 * mm
    _text(canvas, left, y, "以下の通りご請求いたします。", 10.5)
    y -= 16 * mm
    _text(canvas, left, y, "ご請求金額（税込）", 10.5, MUTED)
    _right(canvas, right, y - 1 * mm, f"¥{row.amount_yen:,}", 22)
    canvas.setStrokeColor(INK)
    canvas.line(left, y - 6 * mm, right, y - 6 * mm)

    y -= 28 * mm
    canvas.setFillColor(PALE)
    canvas.rect(left, y - 6 * mm, right - left, 12 * mm, stroke=0, fill=1)
    _text(canvas, left + 4 * mm, y - 1 * mm, "内容", 9.5)
    _right(canvas, right - 4 * mm, y - 1 * mm, "金額（税込）", 9.5)
    y -= 15 * mm
    _text(canvas, left + 4 * mm, y, row.description, 10.5)
    _right(canvas, right - 4 * mm, y, f"¥{row.amount_yen:,}", 10.5)
    canvas.setStrokeColor(LINE)
    canvas.line(left, y - 6 * mm, right, y - 6 * mm)

    y -= 22 * mm
    _text(canvas, left, y, "発行者：［名称・住所・登録番号を確認後に記入］", 10.5)
    y -= 9 * mm
    _text(canvas, left, y, "支払期限・振込先：［運用確認後に記入］", 10.5)
    y -= 14 * mm
    _text(canvas, left, y, "このPDFは機能検証用のサンプルです。税区分・請求先・発行者情報を確認してください。", 8.5, MUTED)
    _right(canvas, right, 22 * mm, f"{row.customer_id}  /  {month}", 8.5, MUTED)
    canvas.save()


def render_eml(path: Path, month: str, row: InvoiceRow, pdf_path: Path) -> None:
    message = EmailMessage()
    message["To"] = row.email
    message["Subject"] = f"【サンプル】{month}分 請求書のご確認"
    message.set_content(
        f"{row.customer_name} 御中\n\n"
        f"{month}分の請求書サンプルを添付します。\n"
        f"金額（税込）：{row.amount_yen:,}円\n\n"
        "送信前に宛先・金額・発行者情報・PDFの内容を確認してください。\n"
        "このファイルは下書きであり、送信処理は行っていません。\n",
        charset="utf-8",
    )
    message.add_attachment(pdf_path.read_bytes(), maintype="application", subtype="pdf", filename=pdf_path.name)
    path.write_bytes(message.as_bytes())


def generate(input_path: Path, month: str, expected_count: int, output_root: Path) -> Path:
    validate_month(month)
    if expected_count < 1:
        raise InvoiceError("expected-count must be at least 1")
    rows = read_rows(input_path)
    if len(rows) != expected_count:
        raise InvoiceError(f"count mismatch: expected {expected_count}, found {len(rows)}; no files generated")
    validate_pdf_fit(rows, month)
    run_name = "run-" + datetime.now().strftime("%Y%m%dT%H%M%S") + "-" + uuid.uuid4().hex[:8]
    run_dir = output_root / month / run_name
    run_dir.mkdir(parents=True, exist_ok=False)
    marker = run_dir / "INCOMPLETE.txt"
    marker.write_text("Generation is incomplete. Do not use files in this folder.\n", encoding="utf-8")
    manifest = run_dir / "manifest.csv"
    pdf_count = eml_count = 0
    try:
        with manifest.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(("customer_id", "customer_name", "email", "amount_yen", "pdf", "eml", "status"))
            for row in rows:
                pdf_name = f"invoice_{month}_{row.customer_id}.pdf"
                eml_name = f"draft_{month}_{row.customer_id}.eml"
                render_pdf(run_dir / pdf_name, month, row)
                pdf_count += 1
                render_eml(run_dir / eml_name, month, row, run_dir / pdf_name)
                eml_count += 1
                writer.writerow((row.customer_id, row.customer_name, row.email, row.amount_yen, pdf_name, eml_name, "draft_created_unsent"))
                stream.flush()
        summary = {
            "month": month,
            "expected_count": expected_count,
            "pdf_count": pdf_count,
            "eml_count": eml_count,
            "total_amount_yen": sum(row.amount_yen for row in rows),
            "status": "complete_unsent",
        }
        (run_dir / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        marker.write_text("Generation complete. Files are unsent drafts requiring human review.\n", encoding="utf-8")
        marker.rename(run_dir / "COMPLETE.txt")
    except Exception as exc:
        raise InvoiceRunError(run_dir) from exc
    return run_dir
